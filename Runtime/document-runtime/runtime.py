#!/usr/bin/env python3
"""Oktapod document runtime v1.

The process accepts newline-delimited JSON-RPC on stdio. Every file argument is
an opaque path relative to OKTAPOD_DOCUMENT_ROOT; project paths, credentials,
shell commands, and URLs are never part of this protocol.
"""

from __future__ import annotations

import hashlib
import io
import json
import os
from pathlib import Path, PurePosixPath
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import traceback
import zipfile

PROTOCOL = "oktapod.document/v1"
VERSION = "1.0.0"
ROOT_VALUE = os.environ.get("OKTAPOD_DOCUMENT_ROOT", "").strip()
ROOT = Path(ROOT_VALUE).resolve() if ROOT_VALUE else None
MAX_INPUT = 100 * 1024 * 1024
MAX_ZIP_ENTRIES = 20_000
MAX_ZIP_EXPANDED = 512 * 1024 * 1024
SUPPORTED = {"pptx", "pdf", "docx", "xlsx"}


class RuntimeFailure(Exception):
    def __init__(self, code: str, message: str, details: dict | None = None):
        super().__init__(message)
        self.code = code
        self.details = details or {}


class _DeniedSocket:
    def __init__(self, *_args, **_kwargs):
        raise RuntimeFailure("NETWORK_DENIED", "document runtime network access is disabled")


# Defense in depth. The Go broker provides only local handles and the child is
# launched with a minimized environment; Python network APIs are disabled too.
socket.socket = _DeniedSocket  # type: ignore[assignment]
socket.create_connection = lambda *_a, **_k: (_ for _ in ()).throw(RuntimeFailure("NETWORK_DENIED", "document runtime network access is disabled"))  # type: ignore[assignment]


def _handle(value: str, *, must_exist: bool = False) -> Path:
    if not ROOT:
        raise RuntimeFailure("RUNTIME_ROOT_MISSING", "OKTAPOD_DOCUMENT_ROOT is required")
    raw = str(value or "").strip().replace("\\", "/")
    pure = PurePosixPath(raw)
    if not raw or pure.is_absolute() or ".." in pure.parts or ":" in raw:
        raise RuntimeFailure("HANDLE_INVALID", "runtime handle is not a safe relative path")
    target = (ROOT / Path(*pure.parts)).resolve()
    if target != ROOT and ROOT not in target.parents:
        raise RuntimeFailure("HANDLE_INVALID", "runtime handle escapes the runtime root")
    if must_exist and (not target.is_file() or target.is_symlink()):
        raise RuntimeFailure("HANDLE_NOT_FOUND", "runtime input handle is not a regular file")
    return target


def _ensure_parent(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.parent.is_symlink():
        raise RuntimeFailure("HANDLE_INVALID", "runtime output parent is a symlink")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return "sha256:" + digest.hexdigest()


def _safe_zip(path: Path, expected: str) -> list[str]:
    warnings: list[str] = []
    try:
        archive = zipfile.ZipFile(path)
    except (zipfile.BadZipFile, OSError) as exc:
        raise RuntimeFailure("OOXML_ZIP_INVALID", f"cannot open OOXML ZIP: {exc}") from exc
    with archive:
        infos = archive.infolist()
        if len(infos) > MAX_ZIP_ENTRIES:
            raise RuntimeFailure("OOXML_ZIP_BOMB", "OOXML entry count exceeds the safety limit")
        expanded = 0
        names = set()
        for info in infos:
            normalized = PurePosixPath(info.filename.replace("\\", "/"))
            if normalized.is_absolute() or ".." in normalized.parts or ":" in info.filename:
                raise RuntimeFailure("OOXML_PATH_TRAVERSAL", "OOXML contains an unsafe entry path")
            if info.flag_bits & 0x1:
                raise RuntimeFailure("DOCUMENT_ENCRYPTED", "encrypted Office files are not supported")
            expanded += info.file_size
            if expanded > MAX_ZIP_EXPANDED or (info.compress_size > 0 and info.file_size / info.compress_size > 2000):
                raise RuntimeFailure("OOXML_ZIP_BOMB", "OOXML expanded size or compression ratio exceeds the safety limit")
            names.add(info.filename)
            if info.filename.lower().endswith(('.xml', '.rels')) and info.file_size <= 8 * 1024 * 1024:
                body = archive.read(info)
                upper = body.upper()
                if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
                    raise RuntimeFailure("OOXML_XML_ENTITY_DENIED", "OOXML contains a DTD or entity declaration")
                if info.filename.lower().endswith('.rels') and b'TARGETMODE="EXTERNAL"' in upper:
                    warnings.append("external_link_preserved_not_loaded")
        required = {
            "pptx": {"[Content_Types].xml", "ppt/presentation.xml"},
            "docx": {"[Content_Types].xml", "word/document.xml"},
            "xlsx": {"[Content_Types].xml", "xl/workbook.xml"},
        }[expected]
        missing = sorted(required - names)
        if missing:
            raise RuntimeFailure("OOXML_PART_MISSING", "required OOXML parts are missing", {"missing": missing})
        if any(name.lower().endswith(("vbaproject.bin", ".xla", ".ole")) for name in names):
            raise RuntimeFailure("DOCUMENT_MACRO_UNSUPPORTED", "macros and embedded executable objects are not supported")
    return sorted(set(warnings))


def _format(path: Path, requested: str | None = None) -> str:
    value = (requested or path.suffix.lstrip(".")).lower()
    if value not in SUPPORTED:
        raise RuntimeFailure("DOCUMENT_FORMAT_UNSUPPORTED", "only .pptx, .pdf, .docx, and .xlsx are supported")
    if path.stat().st_size > MAX_INPUT:
        raise RuntimeFailure("DOCUMENT_TOO_LARGE", "document exceeds the 100 MiB runtime limit")
    if value == "pdf":
        with path.open("rb") as stream:
            if stream.read(5) != b"%PDF-":
                raise RuntimeFailure("PDF_SIGNATURE_INVALID", "PDF does not start with %PDF-")
    else:
        _safe_zip(path, value)
    return value


def _open_pptx(path: Path, mode: str) -> dict:
    from pptx import Presentation
    deck = Presentation(str(path))
    units = []
    for index, slide in enumerate(deck.slides, 1):
        texts, tables, charts, images = [], [], [], 0
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = "\n".join(paragraph.text for paragraph in shape.text_frame.paragraphs).strip()
                if text:
                    texts.append(text)
            if getattr(shape, "has_table", False):
                tables.append([[cell.text for cell in row.cells] for row in shape.table.rows])
            if getattr(shape, "has_chart", False):
                charts.append({"chart_type": str(shape.chart.chart_type), "series_count": len(shape.chart.series)})
            if getattr(shape, "shape_type", None) == 13:
                images += 1
        notes = ""
        try:
            notes = "\n".join(shape.text for shape in slide.notes_slide.notes_text_frame.paragraphs).strip()
        except Exception:
            pass
        units.append({"index": index, "kind": "slide", "text": "\n".join(texts), "notes": notes, "tables": tables if mode == "full" else [], "charts": charts, "image_count": images, "layout": slide.slide_layout.name or ""})
    return {"format": "pptx", "title": deck.core_properties.title or "", "unit_kind": "slide", "unit_count": len(units), "units": units}


def _open_docx(path: Path, mode: str) -> dict:
    from docx import Document
    document = Document(str(path))
    units = []
    for index, paragraph in enumerate(document.paragraphs, 1):
        if paragraph.text or mode == "full":
            units.append({"index": index, "kind": "paragraph", "text": paragraph.text, "style": paragraph.style.name if paragraph.style else "", "runs": [{"text": run.text, "bold": bool(run.bold), "italic": bool(run.italic)} for run in paragraph.runs] if mode == "full" else []})
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    sections = []
    for index, section in enumerate(document.sections, 1):
        sections.append({"index": index, "page_width": int(section.page_width or 0), "page_height": int(section.page_height or 0), "header": "\n".join(p.text for p in section.header.paragraphs), "footer": "\n".join(p.text for p in section.footer.paragraphs)})
    return {"format": "docx", "title": document.core_properties.title or "", "unit_kind": "paragraph", "unit_count": len(units), "units": units, "tables": tables, "sections": sections}


def _open_xlsx(path: Path, mode: str) -> dict:
    from openpyxl import load_workbook
    # Normal mode is required for table, style, and chart metadata. The Go
    # broker and ZIP guard already cap both compressed and expanded input size.
    workbook = load_workbook(str(path), read_only=False, data_only=False, keep_links=False)
    units = []
    try:
        for index, sheet in enumerate(workbook.worksheets, 1):
            rows = []
            max_rows = min(sheet.max_row or 0, 500 if mode == "full" else 100)
            max_columns = min(sheet.max_column or 0, 100)
            for row in sheet.iter_rows(min_row=1, max_row=max_rows, max_col=max_columns):
                rows.append([cell.value for cell in row])
            units.append({"index": index, "kind": "sheet", "name": sheet.title, "max_row": sheet.max_row, "max_column": sheet.max_column, "rows": rows, "truncated": (sheet.max_row or 0) > max_rows or (sheet.max_column or 0) > max_columns, "table_names": sorted(sheet.tables.keys()), "chart_count": len(sheet._charts)})
    finally:
        workbook.close()
    return {"format": "xlsx", "unit_kind": "sheet", "unit_count": len(units), "units": units}


def _tesseract() -> str | None:
    configured = os.environ.get("OKTAPOD_TESSERACT", "").strip()
    if configured and Path(configured).is_file():
        return configured
    return shutil.which("tesseract")


def _ocr_languages() -> list[str]:
    executable = _tesseract()
    if not executable:
        return []
    try:
        completed = subprocess.run(
            [executable, "--list-langs"], stdin=subprocess.DEVNULL,
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=15, check=False,
        )
    except (OSError, subprocess.SubprocessError):
        return []
    if completed.returncode != 0:
        return []
    lines = completed.stdout.decode("utf-8", "replace").splitlines()
    return sorted({line.strip() for line in lines if line.strip() and "available languages" not in line.lower()})


def _ocr_page(pdf_path: Path, page_index: int, work_dir: Path) -> tuple[str, float]:
    executable = _tesseract()
    if not executable:
        return "", 0.0
    import pypdfium2 as pdfium
    document = pdfium.PdfDocument(str(pdf_path))
    try:
        image = document[page_index].render(scale=2).to_pil()
        image_path = work_dir / f"ocr-{page_index + 1}.png"
        image.save(image_path)
    finally:
        document.close()
    completed = subprocess.run([executable, str(image_path), "stdout", "-l", "chi_sim+eng", "tsv"], stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=90, check=False)
    if completed.returncode != 0:
        return "", 0.0
    lines, confidences, words = completed.stdout.decode("utf-8", "replace").splitlines(), [], []
    for line in lines[1:]:
        fields = line.split("\t")
        if len(fields) >= 12 and fields[11].strip():
            words.append(fields[11].strip())
            try:
                confidence = float(fields[10])
                if confidence >= 0:
                    confidences.append(confidence)
            except ValueError:
                pass
    return " ".join(words), (sum(confidences) / len(confidences) / 100 if confidences else 0.0)


def _open_pdf(path: Path, mode: str) -> tuple[dict, list[str]]:
    import pdfplumber
    from pypdf import PdfReader
    warnings: list[str] = []
    reader = PdfReader(str(path), strict=True)
    if reader.is_encrypted:
        raise RuntimeFailure("DOCUMENT_ENCRYPTED", "encrypted PDF files are not supported")
    units = []
    with pdfplumber.open(str(path)) as document:
        for index, page in enumerate(document.pages, 1):
            text = page.extract_text(layout=True) or ""
            tables = page.extract_tables() if mode == "full" else []
            native = len(text.strip()) >= 20
            confidence = 1.0 if native else 0.0
            source = "native"
            if not native and mode == "full" and index <= 20:
                ocr_text, confidence = _ocr_page(path, index - 1, path.parent)
                if ocr_text:
                    text, source = ocr_text, "ocr"
                    warnings.append(f"page_{index}_ocr_confidence_{confidence:.2f}")
                else:
                    warnings.append(f"page_{index}_low_text_density")
            units.append({"index": index, "kind": "page", "text": text, "tables": tables, "width": page.width, "height": page.height, "text_source": source, "confidence": confidence})
    return {"format": "pdf", "unit_kind": "page", "unit_count": len(units), "units": units, "form_fields": sorted((reader.get_fields() or {}).keys())}, warnings


def open_document(params: dict) -> dict:
    path = _handle(params.get("input_handle", ""), must_exist=True)
    mode = params.get("mode", "structure")
    if mode not in {"structure", "full"}:
        raise RuntimeFailure("MODE_INVALID", "mode must be structure or full")
    document_format = _format(path, params.get("format"))
    warnings = _safe_zip(path, document_format) if document_format != "pdf" else []
    if document_format == "pptx":
        structure = _open_pptx(path, mode)
    elif document_format == "docx":
        structure = _open_docx(path, mode)
    elif document_format == "xlsx":
        structure = _open_xlsx(path, mode)
    else:
        structure, pdf_warnings = _open_pdf(path, mode)
        warnings.extend(pdf_warnings)
    return {"format": document_format, "structure": structure, "warnings": sorted(set(warnings)), "content_hash": _sha256(path), "size_bytes": path.stat().st_size}


def _pdf_text_blocks(specification: dict) -> list[dict]:
    """Normalize the supported PDF specification shapes into visible text blocks."""
    blocks: list[dict] = []

    def add(kind: str, value, level: int = 0) -> None:
        text = str(value or "").strip()
        if text:
            blocks.append({"kind": kind, "text": text, "level": level})

    def add_items(items) -> None:
        if isinstance(items, str):
            add("paragraph", items)
            return
        if not isinstance(items, list):
            return
        for item in items:
            if isinstance(item, str):
                add("paragraph", item)
                continue
            if not isinstance(item, dict):
                continue
            kind = str(item.get("kind") or item.get("type") or "paragraph").lower()
            if kind in {"heading", "title"}:
                add("heading", item.get("text") or item.get("title"), int(item.get("level") or 1))
            elif kind in {"bullets", "bullet_list", "list"}:
                for bullet in item.get("items") or item.get("bullets") or []:
                    add("bullet", bullet.get("text") if isinstance(bullet, dict) else bullet)
            else:
                add("paragraph", item.get("text") or item.get("body") or item.get("value"))

    add("title", specification.get("title"))
    add("subtitle", specification.get("subtitle"))
    sections = specification.get("sections")
    if isinstance(sections, list):
        for section in sections:
            if not isinstance(section, dict):
                continue
            add("heading", section.get("heading") or section.get("title"), 1)
            body = section.get("body") or section.get("paragraphs") or []
            add_items(body if isinstance(body, list) else [body])
            for bullet in section.get("bullets") or []:
                add("bullet", bullet.get("text") if isinstance(bullet, dict) else bullet)
    for key in ("content", "blocks", "elements", "body", "items"):
        add_items(specification.get(key))
    nested = specification.get("document")
    if isinstance(nested, dict):
        add("heading", nested.get("title"), 1)
        add_items(nested.get("paragraphs"))
    markdown = specification.get("markdown")
    if isinstance(markdown, str):
        for raw_line in markdown.splitlines():
            line = raw_line.strip()
            if not line:
                continue
            if line.startswith("#"):
                heading = line.lstrip("#").strip()
                add("heading", heading, max(1, min(3, len(line) - len(line.lstrip("#")))))
            elif line.startswith(("- ", "* ", "+ ")):
                add("bullet", line[2:])
            else:
                add("paragraph", line)
    if not blocks:
        add("paragraph", specification.get("text"))
    footer = str(specification.get("footer") or "").strip()
    if footer:
        blocks.append({"kind": "footer", "text": footer, "level": 0})
    return blocks


def _pdf_wrap(text: str, font: str, size: float, width: float) -> list[str]:
    from reportlab.pdfbase.pdfmetrics import stringWidth
    lines: list[str] = []
    for paragraph in str(text).splitlines() or [""]:
        current = ""
        for character in paragraph:
            candidate = current + character
            if current and stringWidth(candidate, font, size) > width:
                lines.append(current.rstrip())
                current = character.lstrip()
            else:
                current = candidate
        lines.append(current.rstrip())
    return lines or [""]


def _create_pdf(output: Path, specification: dict) -> None:
    from reportlab.lib.pagesizes import A4, LETTER
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen.canvas import Canvas

    page_size = LETTER if str(specification.get("page_size", "A4")).lower() == "letter" else A4
    font = "STSong-Light"
    pdfmetrics.registerFont(UnicodeCIDFont(font))
    canvas = Canvas(str(output), pagesize=page_size)
    width, height = page_size
    left, right, top, bottom = 54, 54, 54, 58
    usable_width = width - left - right
    footer_blocks: list[dict] = []

    def draw_footer() -> None:
        if not footer_blocks:
            return
        canvas.setFont(font, 8.5)
        canvas.setFillColorRGB(0.38, 0.42, 0.48)
        footer = " ".join(item["text"] for item in footer_blocks)
        lines = _pdf_wrap(footer, font, 8.5, usable_width)
        canvas.drawString(left, 28, lines[0])

    def new_page() -> float:
        draw_footer()
        canvas.showPage()
        return height - top

    def draw_blocks(blocks: list[dict], *, force_page_after: bool = False) -> None:
        nonlocal_y = height - top
        for block in blocks:
            kind = block["kind"]
            if kind == "footer":
                footer_blocks.append(block)
                continue
            if kind == "title":
                size, leading, before, after = 20, 27, 0, 10
            elif kind == "subtitle":
                size, leading, before, after = 10.5, 16, 0, 12
            elif kind == "heading":
                level = int(block.get("level") or 1)
                size, leading, before, after = (14, 21, 9, 5) if level <= 1 else (12, 18, 7, 4)
            else:
                size, leading, before, after = 10.5, 17, 2, 5
            prefix = "•  " if kind == "bullet" else ""
            text_width = usable_width - (14 if kind == "bullet" else 0)
            lines = _pdf_wrap(prefix + block["text"], font, size, text_width)
            required = before + len(lines) * leading + after
            if nonlocal_y - required < bottom and nonlocal_y < height - top:
                nonlocal_y = new_page()
            nonlocal_y -= before
            canvas.setFont(font, size)
            canvas.setFillColorRGB(0.10, 0.13, 0.18)
            x = left + (14 if kind == "bullet" else 0)
            for line in lines:
                if nonlocal_y < bottom:
                    nonlocal_y = new_page()
                    canvas.setFont(font, size)
                    canvas.setFillColorRGB(0.10, 0.13, 0.18)
                canvas.drawString(x, nonlocal_y, line)
                nonlocal_y -= leading
            nonlocal_y -= after
        if force_page_after:
            new_page()

    pages = specification.get("pages")
    if isinstance(pages, list) and pages:
        for index, page in enumerate(pages):
            if not isinstance(page, dict):
                page = {"text": page}
            page_blocks: list[dict] = []
            title = str(page.get("title") or "").strip()
            if title:
                page_blocks.append({"kind": "heading", "text": title, "level": 1})
            page_blocks.extend(_pdf_text_blocks({key: value for key, value in page.items() if key != "title"}))
            page_blocks = [block for block in page_blocks if block["kind"] != "footer"]
            if not page_blocks:
                raise RuntimeFailure("DOCUMENT_SPECIFICATION_EMPTY", f"PDF page {index + 1} contains no visible text")
            draw_blocks(page_blocks, force_page_after=index < len(pages) - 1)
    else:
        blocks = _pdf_text_blocks(specification)
        if not any(block["kind"] != "footer" and block["text"].strip() for block in blocks):
            raise RuntimeFailure("DOCUMENT_SPECIFICATION_EMPTY", "PDF specification contains no visible text")
        # Register the footer before flowing body content so it is present on
        # every page, including pages created by automatic overflow.
        footer_blocks.extend(block for block in blocks if block["kind"] == "footer")
        draw_blocks([block for block in blocks if block["kind"] != "footer"])
    draw_footer()
    canvas.save()


def create_document(params: dict) -> dict:
    output = _handle(params.get("output_handle", ""))
    _ensure_parent(output)
    document_format = str(params.get("format", "")).lower()
    if document_format not in SUPPORTED:
        raise RuntimeFailure("DOCUMENT_FORMAT_UNSUPPORTED", "unsupported create format")
    template = params.get("template_handle")
    if template:
        source = _handle(template, must_exist=True)
        if _format(source) != document_format:
            raise RuntimeFailure("TEMPLATE_FORMAT_MISMATCH", "template format does not match requested format")
        shutil.copyfile(source, output)
    else:
        specification = params.get("specification") or {}
        if document_format == "pptx":
            from pptx import Presentation
            deck = Presentation()
            for item in specification.get("slides", []):
                slide = deck.slides.add_slide(deck.slide_layouts[1] if len(deck.slide_layouts) > 1 else deck.slide_layouts[0])
                if slide.shapes.title:
                    slide.shapes.title.text = str(item.get("title", ""))
                placeholders = [shape for shape in slide.placeholders if shape != slide.shapes.title and getattr(shape, "has_text_frame", False)]
                if placeholders:
                    placeholders[0].text = str(item.get("body", ""))
            if not specification.get("slides"):
                deck.slides.add_slide(deck.slide_layouts[0])
            deck.save(str(output))
        elif document_format == "docx":
            from docx import Document
            document = Document()
            if specification.get("title"):
                document.add_heading(str(specification["title"]), level=0)
            for item in specification.get("paragraphs", []):
                if isinstance(item, dict):
                    document.add_paragraph(str(item.get("text", "")), style=item.get("style") or None)
                else:
                    document.add_paragraph(str(item))
            document.save(str(output))
        elif document_format == "xlsx":
            from openpyxl import Workbook
            workbook = Workbook()
            default = workbook.active
            sheets = specification.get("sheets", [])
            for index, item in enumerate(sheets):
                sheet = default if index == 0 else workbook.create_sheet()
                sheet.title = str(item.get("name") or f"Sheet{index + 1}")[:31]
                for row in item.get("rows", []):
                    sheet.append(list(row))
            workbook.save(str(output))
        else:
            _create_pdf(output, specification)
    _format(output, document_format)
    return {"output_handle": params["output_handle"], "content_hash": _sha256(output), "size_bytes": output.stat().st_size}


def _replace_pptx(deck, find: str, replace: str) -> int:
    changed = 0
    for slide in deck.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if find in run.text:
                            run.text = run.text.replace(find, replace)
                            changed += 1
    return changed


def apply_document(params: dict) -> dict:
    source = _handle(params.get("input_handle", ""), must_exist=True)
    output = _handle(params.get("output_handle", ""))
    _ensure_parent(output)
    document_format = _format(source, params.get("format"))
    operations = params.get("operations") or []
    applied = []
    if document_format == "pptx":
        from pptx import Presentation
        deck = Presentation(str(source))
        for operation in operations:
            op = operation.get("op")
            if op == "replace_text":
                count = _replace_pptx(deck, str(operation.get("find", "")), str(operation.get("replace", "")))
            elif op == "set_text":
                slide = deck.slides[int(operation.get("slide", 1)) - 1]
                shape = slide.shapes[int(operation.get("shape", 1)) - 1]
                if not getattr(shape, "has_text_frame", False):
                    raise RuntimeFailure("OPERATION_TARGET_INVALID", "target shape has no text frame")
                shape.text = str(operation.get("text", "")); count = 1
            elif op == "add_slide":
                slide = deck.slides.add_slide(deck.slide_layouts[min(int(operation.get("layout", 1)), len(deck.slide_layouts) - 1)])
                if slide.shapes.title:
                    slide.shapes.title.text = str(operation.get("title", ""))
                count = 1
            elif op == "delete_slide":
                slide = deck.slides[int(operation.get("slide", 1)) - 1]
                rel_id = slide._element.getparent().getchildren()[int(operation.get("slide", 1)) - 1].rId
                deck.part.drop_rel(rel_id)
                deck.slides._sldIdLst.remove(deck.slides._sldIdLst[int(operation.get("slide", 1)) - 1]); count = 1
            else:
                raise RuntimeFailure("OPERATION_NOT_IMPLEMENTED", f"{op} is not implemented for PPTX")
            applied.append({"op": op, "changed": count})
        deck.save(str(output))
    elif document_format == "docx":
        from docx import Document
        document = Document(str(source))
        for operation in operations:
            op, count = operation.get("op"), 0
            if op == "replace_text":
                find, replace = str(operation.get("find", "")), str(operation.get("replace", ""))
                for paragraph in document.paragraphs:
                    for run in paragraph.runs:
                        if find in run.text:
                            run.text = run.text.replace(find, replace); count += 1
            elif op in {"add_paragraph", "insert_paragraph"}:
                document.add_paragraph(str(operation.get("text", "")), style=operation.get("style") or None); count = 1
            elif op == "set_text":
                paragraph = document.paragraphs[int(operation.get("paragraph", 1)) - 1]
                paragraph.text = str(operation.get("text", "")); count = 1
            elif op in {"set_header", "set_footer"}:
                section = document.sections[int(operation.get("section", 1)) - 1]
                target = section.header if op == "set_header" else section.footer
                target.paragraphs[0].text = str(operation.get("text", "")); count = 1
            else:
                raise RuntimeFailure("OPERATION_NOT_IMPLEMENTED", f"{op} is not implemented for DOCX")
            applied.append({"op": op, "changed": count})
        document.save(str(output))
    elif document_format == "xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(str(source), keep_links=False)
        for operation in operations:
            op, count = operation.get("op"), 0
            if op in {"set_cell", "set_formula"}:
                sheet = workbook[str(operation.get("sheet") or workbook.sheetnames[0])]
                sheet[str(operation.get("cell", "A1"))] = operation.get("value"); count = 1
            elif op == "add_sheet":
                workbook.create_sheet(str(operation.get("name", "Sheet"))[:31]); count = 1
            elif op == "delete_sheet":
                workbook.remove(workbook[str(operation.get("sheet"))]); count = 1
            elif op == "rename_sheet":
                workbook[str(operation.get("sheet"))].title = str(operation.get("name"))[:31]; count = 1
            elif op == "replace_text":
                find, replace = str(operation.get("find", "")), str(operation.get("replace", ""))
                for sheet in workbook.worksheets:
                    for row in sheet.iter_rows():
                        for cell in row:
                            if isinstance(cell.value, str) and find in cell.value:
                                cell.value = cell.value.replace(find, replace); count += 1
            else:
                raise RuntimeFailure("OPERATION_NOT_IMPLEMENTED", f"{op} is not implemented for XLSX")
            applied.append({"op": op, "changed": count})
        workbook.save(str(output))
    else:
        from pypdf import PdfReader, PdfWriter
        reader, writer = PdfReader(str(source), strict=True), PdfWriter()
        pages = list(reader.pages)
        for operation in operations:
            op, count = operation.get("op"), 0
            indices = [int(value) - 1 for value in operation.get("pages", [])] if operation.get("pages") else list(range(len(pages)))
            if op == "rotate_pages":
                for index in indices:
                    pages[index].rotate(int(operation.get("degrees", 90))); count += 1
            elif op == "reorder_pages":
                order = [int(value) - 1 for value in operation.get("order", [])]
                if sorted(order) != list(range(len(pages))):
                    raise RuntimeFailure("OPERATION_TARGET_INVALID", "page order must contain every page exactly once")
                pages = [pages[index] for index in order]; count = len(pages)
            elif op == "fill_form":
                writer.clone_document_from_reader(reader)
                for page in writer.pages:
                    writer.update_page_form_field_values(page, operation.get("fields") or {}, auto_regenerate=False)
                pages = []; count = len(writer.pages)
            elif op == "crop_pages":
                box = operation.get("box") or {}
                for index in indices:
                    pages[index].cropbox.lower_left = (float(box.get("left", 0)), float(box.get("bottom", 0)))
                    pages[index].cropbox.upper_right = (float(box.get("right", pages[index].mediabox.right)), float(box.get("top", pages[index].mediabox.top))); count += 1
            else:
                raise RuntimeFailure("OPERATION_NOT_IMPLEMENTED", f"{op} is not implemented for PDF")
            applied.append({"op": op, "changed": count})
        for page in pages:
            writer.add_page(page)
        with output.open("wb") as stream:
            writer.write(stream)
    _format(output, document_format)
    return {"output_handle": params["output_handle"], "content_hash": _sha256(output), "size_bytes": output.stat().st_size, "applied": applied}


def _libreoffice() -> str | None:
    configured = os.environ.get("OKTAPOD_LIBREOFFICE", "").strip()
    if configured and Path(configured).is_file():
        # The Windows GUI launcher detaches immediately. Prefer the sibling
        # console launcher so subprocess.run observes conversion completion.
        if os.name == "nt" and Path(configured).name.lower() == "soffice.exe":
            console = Path(configured).with_suffix(".com")
            if console.is_file():
                return str(console)
        return configured
    return shutil.which("soffice") or shutil.which("libreoffice")


def _to_pdf(source: Path, output_dir: Path) -> Path:
    executable = _libreoffice()
    if not executable:
        raise RuntimeFailure("RENDERER_UNAVAILABLE", "LibreOffice Headless is not installed in the document runtime package")
    output_dir.mkdir(parents=True, exist_ok=True)
    # LibreOffice creates deeply nested profile files and can crash on Windows
    # MAX_PATH even when Python can access the requested cache directory. Stage
    # conversion under the runtime's short temp root, independent of cache keys
    # and user filenames. Each call owns its profile and removes it on failure.
    temp_root = ROOT / "tmp"
    temp_root.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="lo-", dir=temp_root) as temporary:
        conversion_dir = Path(temporary)
        profile = conversion_dir / "profile"
        staged_source = conversion_dir / ("input" + source.suffix)
        shutil.copyfile(source, staged_source)
        command = [executable, "--headless", "--nologo", "--nodefault", "--nofirststartwizard", "--norestore", f"-env:UserInstallation={profile.as_uri()}", "--convert-to", "pdf", "--outdir", str(conversion_dir), str(staged_source)]
        completed = subprocess.run(command, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=180, check=False)
        candidate = conversion_dir / "input.pdf"
        if completed.returncode != 0 or not candidate.is_file():
            raise RuntimeFailure("RENDER_FAILED", "LibreOffice conversion failed", {"stderr": completed.stderr.decode("utf-8", "replace")[-2000:]})
        _format(candidate, "pdf")
        result = output_dir / (source.stem + ".pdf")
        shutil.copyfile(candidate, result)
        return result


def render_document(params: dict) -> dict:
    source = _handle(params.get("input_handle", ""), must_exist=True)
    output_dir = _handle(params.get("output_dir_handle", ""))
    output_dir.mkdir(parents=True, exist_ok=True)
    document_format = _format(source, params.get("format"))
    pdf = source if document_format == "pdf" else _to_pdf(source, output_dir)
    import pypdfium2 as pdfium
    document = pdfium.PdfDocument(str(pdf))
    previews = []
    try:
        for index in range(len(document)):
            target = output_dir / f"page-{index + 1:04d}.png"
            document[index].render(scale=1.25).to_pil().save(target)
            previews.append(str(target.relative_to(ROOT)).replace("\\", "/"))
    finally:
        document.close()
    if not previews:
        raise RuntimeFailure("RENDER_FAILED", "renderer produced no pages")
    manifest = {"validated": True, "page_count": len(previews), "previews": previews, "renderer": "PDFium", "rendered_at": int(time.time())}
    body = json.dumps(manifest, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    return {"manifest": manifest, "render_hash": "sha256:" + hashlib.sha256(body).hexdigest(), "warnings": []}


def export_document(params: dict) -> dict:
    source = _handle(params.get("input_handle", ""), must_exist=True)
    output = _handle(params.get("output_handle", ""))
    _ensure_parent(output)
    source_format = _format(source, params.get("source_format"))
    target_format = str(params.get("format", "")).lower()
    if target_format == source_format:
        shutil.copyfile(source, output)
    elif target_format == "pdf" and source_format in {"pptx", "docx", "xlsx"}:
        converted = _to_pdf(source, output.parent)
        if converted != output:
            shutil.move(converted, output)
    else:
        raise RuntimeFailure("EXPORT_FORMAT_UNSUPPORTED", f"cannot export {source_format} as {target_format}")
    _format(output, target_format)
    return {"output_handle": params["output_handle"], "content_hash": _sha256(output), "size_bytes": output.stat().st_size, "format": target_format}


def health(_params: dict) -> dict:
    packages = {}
    for module in ("pptx", "docx", "openpyxl", "pypdf", "pdfplumber", "pypdfium2", "reportlab"):
        try:
            __import__(module); packages[module] = True
        except Exception:
            packages[module] = False
    return {"available": all(packages.values()), "version": VERSION, "protocol": PROTOCOL, "renderer": "LibreOffice+PDFium" if _libreoffice() else "PDFium", "ocr_languages": _ocr_languages(), "packages": packages, "libreoffice": bool(_libreoffice())}


METHODS = {"health": health, "open": open_document, "create": create_document, "apply": apply_document, "render": render_document, "export": export_document}


def main() -> int:
    if not ROOT:
        print(json.dumps({"fatal": "OKTAPOD_DOCUMENT_ROOT is required"}), file=sys.stderr)
        return 2
    ROOT.mkdir(parents=True, exist_ok=True)
    for line in sys.stdin:
        request_id = None
        try:
            request = json.loads(line)
            request_id = request.get("id")
            if request.get("jsonrpc") != "2.0" or request.get("method") not in METHODS or not isinstance(request.get("params", {}), dict):
                raise RuntimeFailure("RPC_REQUEST_INVALID", "invalid JSON-RPC request")
            result = METHODS[request["method"]](request.get("params") or {})
            response = {"jsonrpc": "2.0", "id": request_id, "result": result}
        except RuntimeFailure as exc:
            response = {"jsonrpc": "2.0", "id": request_id, "error": {"code": exc.code, "message": str(exc), "data": exc.details}}
        except Exception as exc:
            response = {"jsonrpc": "2.0", "id": request_id, "error": {"code": "RUNTIME_INTERNAL", "message": str(exc), "data": {"trace": traceback.format_exc(limit=4)}}}
        sys.stdout.write(json.dumps(response, ensure_ascii=False, separators=(",", ":")) + "\n")
        sys.stdout.flush()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
